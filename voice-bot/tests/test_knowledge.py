"""Knowledge (wave D): reading documents and pages, passages, and search in Hindi and English.

The search tests use the real model (multilingual-e5-small). Fetch it once with
`python -m app.fetch_model models/e5` (or point VOICE_BOT_MODEL_DIR at a copy); without it they
are skipped.
"""

import asyncio
import io
import json
import os
import time

import pytest
from fastapi.testclient import TestClient

from app import knowledge
from app.conversation import Conversation
from app.crm import sign, verify

HAS_MODEL = os.path.exists(os.path.join(knowledge.MODEL_DIR, "model_quantized.onnx"))
needs_model = pytest.mark.skipif(not HAS_MODEL, reason="the knowledge model isn't downloaded (python -m app.fetch_model)")
SECRET = "test-secret-0123456789abcdef0123456789"

FACTS = [
    ("Fees", "The MBA programme fee is 2,50,000 rupees per year, payable in two instalments."),
    ("Intakes", "The next MBA intakes start in January 2027 and July 2027. Applications close a month before each start."),
    ("Eligibility", "Eligibility: a bachelor's degree with at least 50 percent marks in any subject."),
    ("Hostel", "Hostel rooms are available on campus for 60,000 rupees per year, including meals."),
    ("Placements", "Placements: over 120 companies visited last year; the average package was 8 lakh rupees."),
    ("Scholarships", "छात्रवृत्ति: 85 प्रतिशत से अधिक अंक वाले छात्रों को फीस में 20 प्रतिशत की छूट मिलती है।"),
]


# ---- Passages and reading ------------------------------------------------------------------------
def test_passages_follow_paragraphs_and_stay_short():
    text = "\n\n".join(["Short first paragraph."] + ["A sentence about the programme and its fees. " * 6] * 8)
    passages = knowledge.split_passages(text)
    assert passages[0].startswith("Short first paragraph.")
    assert all(len(passage) <= knowledge.CHUNK_CHARS for passage in passages)
    assert len(passages) >= 3
    long_line = "word " * 400
    assert all(len(passage) <= knowledge.CHUNK_CHARS for passage in knowledge.split_passages(long_line))
    assert knowledge.split_passages("   \n\n  ") == []


def test_hindi_sentences_split_at_the_danda():
    text = ("यह एक लंबा वाक्य है जो प्रोग्राम के बारे में बताता है। " * 30).strip()
    passages = knowledge.split_passages(text)
    assert len(passages) > 1 and all(passage.endswith("।") for passage in passages)


def test_reads_a_web_page_without_menus_and_scripts():
    html = """<html><head><title>MBA Admissions</title><script>var x = 1;</script></head><body>
    <nav><a href="/">Home</a><a href="/about">About</a></nav>
    <main><h1>MBA</h1><p>The fee is 2,50,000 rupees a year.</p><ul><li>January intake</li><li>July intake</li></ul></main>
    <footer>Copyright 2026</footer></body></html>"""
    title, text = knowledge.extract_page(html)
    assert title == "MBA Admissions"
    assert "The fee is 2,50,000 rupees a year." in text and "January intake" in text
    assert "Home" not in text and "Copyright" not in text and "var x" not in text


def test_reads_word_and_text_files():
    from docx import Document

    document = Document()
    document.add_paragraph("Hostel rooms cost 60,000 rupees a year.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Application fee"
    table.rows[0].cells[1].text = "1,000"
    buffer = io.BytesIO()
    document.save(buffer)
    text = knowledge.extract_document("docx", buffer.getvalue())
    assert "Hostel rooms cost 60,000 rupees a year." in text and "Application fee | 1,000" in text
    assert knowledge.extract_document("text", "हॉस्टल उपलब्ध है।\r\n\r\n\r\nYes.".encode()) == "हॉस्टल उपलब्ध है।\n\nYes."
    with pytest.raises(ValueError):
        knowledge.extract_document("exe", b"MZ")


def test_vectors_round_trip():
    import numpy as np

    vector = np.arange(knowledge.DIMENSIONS, dtype=np.float32) / 1000
    assert np.allclose(knowledge.decode_vector(knowledge.encode_vector(vector)), vector)


# ---- Search ------------------------------------------------------------------------------------------
@pytest.fixture(scope="module")
def index():
    embedder = knowledge.Embedder.get()
    vectors = embedder.embed([text for _, text in FACTS], "passage")
    items = [{"text": text, "source": source, "embedding": knowledge.encode_vector(vector)} for (source, text), vector in zip(FACTS, vectors)]
    return knowledge.KnowledgeIndex(items)


@needs_model
@pytest.mark.parametrize("question,source", [
    # English, as the AI writes its search (the tool asks for English).
    ("MBA fee", "Fees"),
    ("When does the next MBA batch start", "Intakes"),
    ("Who can apply", "Eligibility"),
    ("Is hostel available", "Hostel"),
    ("placement record", "Placements"),
    ("Is there a scholarship", "Scholarships"),  # the fact is written only in Hindi
    # The caller's own words, in case the AI passes them on as they were.
    ("MBA की फीस कितनी है?", "Fees"),
    ("अगला बैच कब शुरू होगा?", "Intakes"),
    ("हॉस्टल की सुविधा है क्या?", "Hostel"),
    ("mba ki fees kitni hai", "Fees"),
    ("kya hostel milega", "Hostel"),
])
def test_finds_the_right_fact_in_hindi_and_english(index, question, source):
    found = index.search(question)
    assert source in [item["source"] for item in found], (question, found)


@needs_model
def test_search_is_fast_and_bounded(index):
    index.search("warm up")
    started = time.perf_counter()
    for _ in range(20):
        found = index.search("What is the fee for the MBA programme?")
    per_question_ms = (time.perf_counter() - started) * 1000 / 20
    assert 1 <= len(found) <= knowledge.TOP_K
    assert per_question_ms < 50, per_question_ms
    print(f"\nKnowledge search: {per_question_ms:.1f} ms a question over {len(index)} passages")


@needs_model
def test_prepare_returns_passages_with_language_and_vectors():
    passages = knowledge.prepare("Q: Is there a hostel?\nA: Yes, for 60,000 rupees a year.", whole=True)
    assert len(passages) == 1 and passages[0]["language"] == "en"
    assert knowledge.decode_vector(passages[0]["embedding"]).shape == (knowledge.DIMENSIONS,)
    assert knowledge.prepare("") == []


@needs_model
def test_the_tool_answers_from_knowledge(index):
    from tests.test_units import Params, bot

    items = [{"text": passage.text, "source": passage.source, "embedding": knowledge.encode_vector(index.matrix[position])} for position, passage in enumerate(index.passages)]
    convo = Conversation({"sessionId": "s", "kind": "TEST", "bot": bot(), "keys": {}, "audio": {}, "knowledge": items}, sample_rate=24000)
    params = Params(arguments={"question": "hostel fees"})
    asyncio.run(convo._tool_search_knowledge(params))
    assert params.result["found"] is True and params.result["passages"][0]["source"] == "Hostel"
    assert convo.events[-1]["type"] == "knowledge_searched" and convo.events[-1]["detail"]["found"] is True
    empty = Conversation({"sessionId": "s", "kind": "TEST", "bot": bot(), "keys": {}, "audio": {}}, sample_rate=24000)
    nothing = Params(arguments={"question": "hostel fees"})
    asyncio.run(empty._tool_search_knowledge(nothing))
    assert nothing.result["found"] is False


# ---- The CRM's signed calls -------------------------------------------------------------------------
def test_signature_check():
    body = '{"text":"x"}'
    header = sign(SECRET, "POST", "/internal/knowledge/prepare", body, 1793000000)
    assert verify(SECRET, header, "POST", "/internal/knowledge/prepare", body, now=1793000030)
    assert not verify(SECRET, header, "POST", "/internal/knowledge/prepare", body, now=1793000061)
    assert not verify(SECRET, header, "POST", "/internal/knowledge/page", body, now=1793000000)
    assert not verify(SECRET, header, "POST", "/internal/knowledge/prepare", '{"text":"y"}', now=1793000000)
    assert not verify(SECRET, "t=1793000000,v1=nothex", "POST", "/internal/knowledge/prepare", body, now=1793000000)
    assert not verify("", header, "POST", "/internal/knowledge/prepare", body, now=1793000000)


def test_internal_endpoints_need_the_signature():
    import app.main as main_module

    client = TestClient(main_module.app)
    body = json.dumps({"html": "<html><body><p>Fees are 1,000 rupees.</p></body></html>"})
    assert client.post("/internal/knowledge/page", content=body).status_code == 404
    signed = client.post("/internal/knowledge/page", content=body, headers={"x-voice-bot-signature": sign(SECRET, "POST", "/internal/knowledge/page", body)})
    assert signed.status_code == 200 and "Fees are 1,000 rupees." in signed.json()["text"]
    bad = json.dumps({"fileType": "docx", "data": "bm90IGEgZG9jeA=="})
    refused = client.post("/internal/knowledge/document", content=bad, headers={"x-voice-bot-signature": sign(SECRET, "POST", "/internal/knowledge/document", bad)})
    assert refused.status_code == 422
