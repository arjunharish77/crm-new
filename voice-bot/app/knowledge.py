"""The bot's knowledge (voice bot plan wave D, decision 5).

The CRM keeps the sources; this service does the work that needs Python:
- reads text out of PDF, Word and web pages;
- splits it into short passages;
- turns passages and questions into vectors with multilingual-e5-small (on the CPU, a few
  milliseconds a question);
- during a call, finds the passages closest to what the caller asked.

The search runs in memory over the bot's passages, which the CRM hands over with the call.
"""

import base64
import io
import os
import re
import threading
from dataclasses import dataclass

import numpy as np

from .language import detect

MODEL_DIR = os.environ.get("VOICE_BOT_MODEL_DIR", os.path.join(os.path.dirname(os.path.dirname(__file__)), "models", "e5"))
DIMENSIONS = 384
CHUNK_CHARS = 700
CHUNK_OVERLAP = 120
MAX_TEXT_CHARS = 400_000
TOP_K = 3
# Passages are returned when they're among the closest and not far behind the best one. e5's
# scores sit close together (0.75 to 0.92), so the AI, told to use only what actually answers,
# makes the last call.
SCORE_FLOOR = 0.75
SCORE_SPREAD = 0.06


class Embedder:
    """multilingual-e5-small, loaded once on first use."""

    _lock = threading.Lock()
    _instance: "Embedder | None" = None

    def __init__(self, folder: str):
        import onnxruntime as ort
        from tokenizers import Tokenizer

        self._tokenizer = Tokenizer.from_file(os.path.join(folder, "tokenizer.json"))
        self._tokenizer.enable_truncation(max_length=512)
        self._tokenizer.enable_padding()
        options = ort.SessionOptions()
        options.intra_op_num_threads = 1
        self._session = ort.InferenceSession(os.path.join(folder, "model_quantized.onnx"), options, providers=["CPUExecutionProvider"])
        self._inputs = {item.name for item in self._session.get_inputs()}

    @classmethod
    def get(cls) -> "Embedder":
        with cls._lock:
            if cls._instance is None:
                cls._instance = Embedder(MODEL_DIR)
            return cls._instance

    def embed(self, texts: list[str], kind: str) -> np.ndarray:
        """Unit vectors, one row per text. `kind` is "query" or "passage" (e5's prefixes)."""
        rows = []
        for start in range(0, len(texts), 32):
            batch = [f"{kind}: {text}" for text in texts[start : start + 32]]
            encoded = self._tokenizer.encode_batch(batch)
            ids = np.array([item.ids for item in encoded], dtype=np.int64)
            mask = np.array([item.attention_mask for item in encoded], dtype=np.int64)
            feeds = {"input_ids": ids, "attention_mask": mask}
            if "token_type_ids" in self._inputs:
                feeds["token_type_ids"] = np.zeros_like(ids)
            hidden = self._session.run(None, feeds)[0]
            pooled = (hidden * mask[..., None]).sum(axis=1) / np.maximum(mask.sum(axis=1, keepdims=True), 1)
            rows.append(pooled / np.linalg.norm(pooled, axis=1, keepdims=True))
        return np.vstack(rows).astype(np.float32) if rows else np.zeros((0, DIMENSIONS), dtype=np.float32)


# ---- Reading documents and pages -------------------------------------------------------------
def tidy_text(text: str) -> str:
    text = text.replace(" ", " ").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()[:MAX_TEXT_CHARS]


def extract_document(file_type: str, data: bytes) -> str:
    """Plain text from a PDF, Word (.docx) or text file."""
    if file_type == "pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return tidy_text("\n\n".join((page.extract_text() or "") for page in reader.pages[:300]))
    if file_type == "docx":
        from docx import Document

        document = Document(io.BytesIO(data))
        parts = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return tidy_text("\n".join(parts))
    if file_type == "text":
        return tidy_text(data.decode("utf-8", errors="replace"))
    raise ValueError("Unsupported file type")


def extract_page(html: str) -> tuple[str, str]:
    """(title, readable text) of a web page, without menus, scripts and footers."""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    title = (soup.title.string or "").strip() if soup.title and soup.title.string else ""
    for tag in soup(["script", "style", "noscript", "svg", "nav", "header", "footer", "form", "iframe", "aside", "button"]):
        tag.decompose()
    main = soup.find("main") or soup.find("article") or soup.body or soup
    blocks = []
    for element in main.find_all(["h1", "h2", "h3", "h4", "p", "li", "td", "th", "dt", "dd", "blockquote"]):
        text = " ".join(element.get_text(" ", strip=True).split())
        if len(text) >= 3 and (not blocks or blocks[-1] != text):
            blocks.append(text)
    if not blocks:
        blocks = [" ".join(main.get_text(" ", strip=True).split())]
    return title[:200], tidy_text("\n".join(blocks))


# ---- Passages --------------------------------------------------------------------------------
SENTENCE_END = re.compile(r"(?<=[.!?।])\s+")


def split_passages(text: str) -> list[str]:
    """Passages of about CHUNK_CHARS, cut at paragraph and sentence ends, with a little overlap
    so a fact split across two passages is still found."""
    text = tidy_text(text)
    if not text:
        return []
    pieces: list[str] = []
    for paragraph in re.split(r"\n\s*\n", text):
        paragraph = paragraph.strip()
        if not paragraph:
            continue
        if len(paragraph) <= CHUNK_CHARS:
            pieces.append(paragraph)
        else:
            pieces.extend(sentence.strip() for sentence in SENTENCE_END.split(paragraph) if sentence.strip())
    passages: list[str] = []
    current = ""
    for piece in pieces:
        while len(piece) > CHUNK_CHARS:  # a very long sentence or line
            cut = piece.rfind(" ", 0, CHUNK_CHARS)
            cut = cut if cut > CHUNK_CHARS // 2 else CHUNK_CHARS
            if current:
                passages.append(current)
                current = ""
            passages.append(piece[:cut].strip())
            piece = piece[cut:].strip()
        if not current:
            current = piece
        elif len(current) + 1 + len(piece) <= CHUNK_CHARS:
            current = f"{current}\n{piece}"
        else:
            passages.append(current)
            tail = current[-CHUNK_OVERLAP:]
            tail = tail[tail.find(" ") + 1 :] if " " in tail else ""
            current = f"{tail} {piece}".strip() if tail else piece
    if current:
        passages.append(current)
    return passages


def encode_vector(vector: np.ndarray) -> str:
    return base64.b64encode(np.asarray(vector, dtype="<f4").tobytes()).decode()


def decode_vector(encoded: str) -> np.ndarray:
    return np.frombuffer(base64.b64decode(encoded), dtype="<f4")


def prepare(text: str, *, whole: bool = False) -> list[dict]:
    """Passages of `text` with their language and vector. `whole` keeps it as one passage (a Q&A)."""
    passages = [tidy_text(text)] if whole else split_passages(text)
    passages = [passage for passage in passages if passage]
    if not passages:
        return []
    vectors = Embedder.get().embed(passages, "passage")
    return [{"text": passage, "language": detect(passage), "embedding": encode_vector(vector)} for passage, vector in zip(passages, vectors)]


# ---- Search during a call -----------------------------------------------------------------------
@dataclass
class Passage:
    text: str
    source: str


class KnowledgeIndex:
    def __init__(self, items: list[dict]):
        self.passages = [Passage(text=str(item.get("text", "")), source=str(item.get("source", ""))) for item in items]
        vectors = [decode_vector(item["embedding"]) for item in items if item.get("embedding")]
        valid = len(vectors) == len(self.passages) and all(vector.shape == (DIMENSIONS,) for vector in vectors)
        self.matrix = np.vstack(vectors) if valid and vectors else np.zeros((0, DIMENSIONS), dtype=np.float32)
        if not valid:
            self.passages = []

    def __len__(self) -> int:
        return len(self.passages)

    def search(self, question: str, embedder: "Embedder | None" = None) -> list[dict]:
        if not self.passages or not question.strip():
            return []
        query = (embedder or Embedder.get()).embed([question], "query")[0]
        scores = self.matrix @ query
        order = np.argsort(-scores)[:TOP_K]
        best = float(scores[order[0]])
        found = []
        for index in order:
            score = float(scores[index])
            if score >= SCORE_FLOOR and score >= best - SCORE_SPREAD:
                found.append({"text": self.passages[index].text, "source": self.passages[index].source, "score": round(score, 3)})
        return found
